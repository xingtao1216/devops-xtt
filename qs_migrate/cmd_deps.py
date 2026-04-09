"""deps 子命令：导出资源依赖关系到 Excel"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from qs_common import log, paginate, ResourceResolver


HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")


def _style_header(ws):
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def _auto_width(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 60)


def export_deps(qs, account_id, resource_type, values, output):
    """导出指定资源或全量资源的依赖关系"""
    resolver = ResourceResolver(qs, account_id)

    ds_map = {d["DataSourceId"]: d["Name"] for d in resolver.datasources()}
    dset_map = {d["DataSetId"]: d["Name"] for d in resolver.datasets()}
    an_map = {d["AnalysisId"]: d["Name"] for d in resolver.analyses()}
    db_map = {d["DashboardId"]: d["Name"] for d in resolver.dashboards()}

    wb = Workbook()

    # ---- Sheet 1: Dashboard 依赖 ----
    ws1 = wb.active
    ws1.title = "Dashboard Dependencies"
    ws1.append(["Dashboard Name", "Dashboard ID", "Analysis Name", "Analysis ID",
                "DataSet Name", "DataSet ID", "DataSource Name", "DataSource ID"])

    dashboards = _get_targets(resolver, "dashboard", values, resource_type, resolver.dashboards, "DashboardId")
    for db in dashboards:
        db_id = db["DashboardId"]
        db_name = db["Name"]
        an_arn = resolver.get_dashboard_analysis(db_id)
        an_id = an_arn.split("/")[-1] if an_arn else ""
        an_name = an_map.get(an_id, "") if an_id else "(无关联 Analysis)"

        ds_arns = resolver.get_dashboard_dataset_arns(db_id)
        if not ds_arns:
            ws1.append([db_name, db_id, an_name, an_id, "(无)", "", "(无)", ""])
            continue
        for ds_arn in sorted(ds_arns):
            did = ds_arn.split("/")[-1]
            dname = dset_map.get(did, "(未知)")
            src_arns = resolver.get_dataset_datasource_arns(did)
            if not src_arns:
                ws1.append([db_name, db_id, an_name, an_id, dname, did, "(无)", ""])
            else:
                for sa in sorted(src_arns):
                    sid = sa.split("/")[-1]
                    sname = ds_map.get(sid, "(未知)")
                    ws1.append([db_name, db_id, an_name, an_id, dname, did, sname, sid])

    _style_header(ws1)
    _auto_width(ws1)

    # ---- Sheet 2: Analysis 依赖 ----
    ws2 = wb.create_sheet("Analysis Dependencies")
    ws2.append(["Analysis Name", "Analysis ID", "DataSet Name", "DataSet ID", "DataSource Name", "DataSource ID"])

    analyses = _get_targets(resolver, "analysis", values, resource_type, resolver.analyses, "AnalysisId")
    for an in analyses:
        an_id = an["AnalysisId"]
        an_name = an["Name"]
        ds_arns = resolver.get_analysis_dataset_arns(an_id)
        if not ds_arns:
            ws2.append([an_name, an_id, "(无)", "", "(无)", ""])
            continue
        for ds_arn in sorted(ds_arns):
            did = ds_arn.split("/")[-1]
            dname = dset_map.get(did, "(未知)")
            src_arns = resolver.get_dataset_datasource_arns(did)
            if not src_arns:
                ws2.append([an_name, an_id, dname, did, "(无)", ""])
            else:
                for sa in sorted(src_arns):
                    sid = sa.split("/")[-1]
                    sname = ds_map.get(sid, "(未知)")
                    ws2.append([an_name, an_id, dname, did, sname, sid])

    _style_header(ws2)
    _auto_width(ws2)

    # ---- Sheet 3: DataSet 依赖 ----
    ws3 = wb.create_sheet("DataSet Dependencies")
    ws3.append(["DataSet Name", "DataSet ID", "ImportMode", "DataSource Name", "DataSource ID"])

    datasets = _get_targets(resolver, "dataset", values, resource_type, resolver.datasets, "DataSetId")
    for ds in datasets:
        did = ds["DataSetId"]
        dname = ds["Name"]
        mode = ds.get("ImportMode", "")
        src_arns = resolver.get_dataset_datasource_arns(did)
        if not src_arns:
            ws3.append([dname, did, mode, "(无)", ""])
        else:
            for sa in sorted(src_arns):
                sid = sa.split("/")[-1]
                sname = ds_map.get(sid, "(未知)")
                ws3.append([dname, did, mode, sname, sid])

    _style_header(ws3)
    _auto_width(ws3)

    wb.save(output)
    log("OK", f"依赖关系已导出: {output}")


def _get_targets(resolver, target_type, values, requested_type, list_fn, id_key):
    """如果用户指定了资源则过滤，否则返回全量"""
    if not values:
        return list_fn()
    if requested_type != target_type:
        return list_fn()
    # 按 id 或 name 过滤
    all_items = list_fn()
    result = []
    for v in values:
        for item in all_items:
            if item[id_key] == v or item.get("Name") == v:
                result.append(item)
    return result if result else all_items
